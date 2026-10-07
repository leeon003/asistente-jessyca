"""Comprehensive Test Suite for ARE Real Tools Workflows (Phase 13).

Validates:
1. Catálogo real de capacidades: parámetros, salidas, límites, verificadores y rechazo de inexistentes.
2. Objetivo 1: Buscar información pública, conservar fuentes y elaborar informe Word (.docx).
3. Objetivo 2: Leer datos de hoja de cálculo, calcular resumen tipado y crear presentación PowerPoint (.pptx).
4. Objetivo 3: Localizar archivos en carpeta, proponer organización y ejecutar cambios autorizados.
5. Propagación de resultados tipados entre pasos ($step.output.prop) y detención estricta de dependientes.
6. Conservación del estado de ejecución y explicación auditada (qué se completó, qué falló, qué puede retomarse).
7. Idempotencia en reanudación/reintento: No repetir escrituras físicas ya verificadas en disco.
8. Solicitud de aclaración ante destinos o recursos esenciales faltantes.
9. Cancelación controlada con CancellationToken e invariante de gobernanza JDE.
10. Variaciones dinámicas de nombres, rutas y datos para evitar acoplamiento a ejemplos fijos.
"""

from __future__ import annotations

import csv
import json
import os
import shutil
import tempfile
from pathlib import Path
import pytest
import docx
import openpyxl
import pptx

from core.are.action_catalog import (
    ActionCatalog,
    ActionVerificationStrategy,
    ExecutableAction,
    get_default_action_catalog,
)
from core.are.dataflow import DataFlowResolver
from core.are.exceptions import (
    MissingOutputDependencyError,
    PlanClarificationRequired,
    PlanValidationError,
    ReplanningDenialViolationError,
    SandboxViolationError,
    UnresolvableActionError,
)
from core.are.execution_state import AREPlanExecutionState
from core.are.models import ARERuntimeMode
from core.are.plan_validator import PlanValidator
from core.are.runtime import ARERuntime, reset_are_runtime
from core.jde.runtime_binding import (
    JDERuntimeBinding,
    JDERuntimeVerdict,
    JDERuntimeVerdictType,
)
from core.emergency_stop import CancellationToken
from core.local_agent.local_agent_models import AgentExecutionState, JessycaRequest
from core.security import RiskLevel


class MockAgentForPhase13:
    """Mock agent governed by JDE for Phase 13 validation."""

    def __init__(self, jde_binding: JDERuntimeBinding | None = None):
        self.jde_runtime_binding = jde_binding or JDERuntimeBinding()
        self.context_manager = None
        self._latest_metrics = None
        self.skill_manager = None
        self.tool_registry = None


# ── TEST 1: CATÁLOGO REAL DE CAPACIDADES Y RECHAZO DE INEXISTENTES ──

def test_catalog_real_capabilities_and_reject_non_existent():
    """Verifica que el catálogo exponga las capacidades reales con parámetros, salidas,
    límites y verificadores, y que rechace estrictamente herramientas inexistentes."""
    catalog = ActionCatalog()

    # Comprobar presencia de acciones para los tres objetivos
    expected_actions = [
        "browser.search",
        "browser.read",
        "documents.create_word_report",
        "word.create_report",
        "spreadsheet.read_sheet",
        "excel.read_data",
        "data.calculate_summary",
        "presentation.create_presentation",
        "powerpoint.create_presentation",
        "files.search",
        "files.preview_organize",
        "files.apply_organization",
        "files.verify_copies",
    ]

    for aid in expected_actions:
        act = catalog.get_action(aid)
        assert act is not None, f"Acción '{aid}' debe estar registrada en ActionCatalog."
        assert act.input_schema is not None
        assert act.output_schema is not None
        assert act.verification_strategy is not None
        assert act.handler is not None

    # Rechazo estricto de herramientas inexistentes y placeholders genéricos
    with pytest.raises(UnresolvableActionError) as exc_generic:
        catalog.resolve_action("action.execute")
    assert "placeholder genérico" in str(exc_generic.value)

    with pytest.raises(UnresolvableActionError) as exc_missing:
        catalog.resolve_action("tools.herramienta_fantasma_inventada")
    assert exc_missing.value.code == "ACTION_NOT_FOUND"


# ── TEST 2: OBJETIVO 1 — BÚSQUEDA PÚBLICA, CONSERVACIÓN DE FUENTES E INFORME WORD ──

@pytest.mark.parametrize(
    "query,doc_name,subfolder",
    [
        ("energias renovables y transicion", "informe_renovables.docx", "reportes_energia"),
        ("ciberseguridad en infraestructuras criticas", "analisis_ciberseguridad.docx", "analisis_ti"),
        ("inteligencia artificial en biomedicina", "estudio_ia_medicina.docx", "biomed"),
    ],
)
def test_objective_1_search_sources_and_create_word_report(query: str, doc_name: str, subfolder: str):
    """Objetivo 1:
    1. Buscar información pública sobre un tema dinámico y conservar fuentes auditables.
    2. Elaborar un informe Word (.docx) estructurado con secciones, contenido y bibliografía de fuentes.
    3. Comprobar físicamente la existencia e integridad del documento en disco."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        p_tmp = Path(tmp_dir)
        target_dir = p_tmp / subfolder
        target_dir.mkdir(parents=True, exist_ok=True)
        docx_path = target_dir / doc_name

        rt = ARERuntime(mode=ARERuntimeMode.CONTROLLED)
        agent = MockAgentForPhase13()

        plan_proposal = {
            "goal": f"Investigar {query} y elaborar un informe Word profesional con fuentes",
            "steps": [
                {
                    "step_id": "step_1_search",
                    "title": f"Búsqueda pública sobre {query}",
                    "action_type": "browser.search",
                    "parameters": {"query": query, "max_results": 3},
                    "depends_on": [],
                },
                {
                    "step_id": "step_2_word",
                    "title": "Elaborar informe Word estructurado con fuentes",
                    "action_type": "documents.create_word_report",
                    "parameters": {
                        "file_path": str(docx_path),
                        "title": f"Informe de Investigación: {query.title()}",
                        "subtitle": "Análisis cuantitativo y fuentes públicas consultadas",
                        "sections": [
                            {
                                "heading": "1. Resumen Ejecutivo",
                                "heading_level": 1,
                                "paragraphs": [
                                    f"Este documento consolida la investigación preliminar sobre {query}.",
                                    "Los hallazgos muestran un crecimiento sostenido en la adopción de nuevas metodologías.",
                                ],
                                "bullets": [
                                    "Identificación de tendencias emergentes.",
                                    "Validación de fuentes primarias auditadas.",
                                ],
                            },
                            {
                                "heading": "2. Indicadores y Métricas Relevantes",
                                "heading_level": 1,
                                "table_headers": ["Dimensión", "Estado", "Nivel de Impacto"],
                                "table_rows": [
                                    ["Madurez Tecnológica", "En progreso", "Alto"],
                                    ["Marco Regulatorio", "Definido", "Medio"],
                                    ["Sostenibilidad", "Verificada", "Muy Alto"],
                                ],
                            },
                        ],
                        "sources": "$step_1_search.output.sources",
                        "author": "JESSYCA 3.0 Core",
                    },
                    "depends_on": ["step_1_search"],
                },
            ],
        }

        req = JessycaRequest(request_id=f"req_obj1_{Path(doc_name).stem}", user_input=f"investiga {query}")
        resp = rt.execute_llm_plan(agent, req, plan_proposal, sandbox_dir=p_tmp)

        assert resp.success is True, f"Error en ejecución del plan: {resp.error}"
        assert resp.status == AgentExecutionState.COMPLETED
        assert len(resp.tools_executed) == 2

        # 1. Verificar salida tipada del paso de búsqueda
        search_out = resp.output_data["step_1_search"]
        assert search_out["verified"] is True
        assert len(search_out["sources"]) >= 1
        assert all("url" in s and "title" in s for s in search_out["sources"])

        # 2. Verificar salida tipada del paso Word
        word_out = resp.output_data["step_2_word"]
        assert word_out["verified"] is True
        assert word_out["size_bytes"] > 0
        assert word_out["sources_count"] == len(search_out["sources"])

        # 3. Comprobar físicamente en disco con python-docx
        assert docx_path.exists(), f"El archivo Word {docx_path} debe existir físicamente."
        assert docx_path.stat().st_size > 0

        doc = docx.Document(str(docx_path))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert query.title() in all_text or "Informe de Investigación" in all_text
        assert "Fuentes Consultadas y Bibliografía" in all_text
        # Comprobar que las fuentes aparecen en las viñetas del documento
        first_url = search_out["sources"][0]["url"]
        assert first_url in all_text
        # Comprobar tabla
        assert len(doc.tables) == 1
        assert len(doc.tables[0].rows) == 4  # 1 header + 3 filas de datos


# ── TEST 3: OBJETIVO 2 — LEER HOJA DE CÁLCULO, CALCULAR RESUMEN Y CREAR PRESENTACIÓN ──

@pytest.mark.parametrize(
    "data_format,ppt_name",
    [
        ("xlsx", "presentacion_presupuesto.pptx"),
        ("csv", "resumen_operaciones.pptx"),
    ],
)
def test_objective_2_spreadsheet_summary_and_powerpoint(data_format: str, ppt_name: str):
    """Objetivo 2:
    1. Leer datos estructurados de una hoja de cálculo (Excel .xlsx o CSV) con valores dinámicos.
    2. Calcular resumen estadístico tipado (totales, promedios, min, max, conteos).
    3. Crear una presentación PowerPoint (.pptx) de 16:9 con diapositivas y tabla consolidada.
    4. Comprobar físicamente los resultados calculados y la presentación en disco."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        p_tmp = Path(tmp_dir)
        source_sheet = p_tmp / f"datos_financieros.{data_format}"
        ppt_path = p_tmp / "presentaciones" / ppt_name
        ppt_path.parent.mkdir(parents=True, exist_ok=True)

        headers = ["Departamento", "Servidores", "Costo_Licencias", "Presupuesto_Total"]
        rows = [
            ["Infraestructura", 10, 4500.0, 15000.0],
            ["Desarrollo", 25, 12000.0, 32000.0],
            ["Seguridad", 5, 8500.0, 18500.0],
            ["Operaciones", 12, 3200.0, 14200.0],
        ]

        if data_format == "xlsx":
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Finanzas"
            ws.append(headers)
            for r in rows:
                ws.append(r)
            wb.save(str(source_sheet))
        else:
            with open(source_sheet, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerows(rows)

        rt = ARERuntime(mode=ARERuntimeMode.CONTROLLED)
        agent = MockAgentForPhase13()

        plan_proposal = {
            "goal": "Analizar hoja de cálculo de finanzas y elaborar presentación ejecutiva",
            "steps": [
                {
                    "step_id": "step_1_read",
                    "title": "Leer hoja de cálculo con datos financieros",
                    "action_type": "spreadsheet.read_sheet",
                    "parameters": {"file_path": str(source_sheet)},
                    "depends_on": [],
                },
                {
                    "step_id": "step_2_calc",
                    "title": "Calcular resumen estadístico cuantitativo tipado",
                    "action_type": "data.calculate_summary",
                    "parameters": {
                        "headers": "$step_1_read.output.headers",
                        "rows": "$step_1_read.output.rows",
                        "numeric_columns": ["Costo_Licencias", "Presupuesto_Total"],
                        "category_column": "Departamento",
                    },
                    "depends_on": ["step_1_read"],
                },
                {
                    "step_id": "step_3_ppt",
                    "title": "Crear presentación ejecutiva PowerPoint con los resultados",
                    "action_type": "presentation.create_presentation",
                    "parameters": {
                        "file_path": str(ppt_path),
                        "title": "Informe Ejecutivo de Finanzas y Presupuesto",
                        "subtitle": "Consolidado de Licencias y Presupuesto por Área",
                        "summary_stats": "$step_2_calc.output.summary_stats",
                        "key_insights": "$step_2_calc.output.key_insights",
                        "presenter": "Equipo de Inteligencia Operativa",
                    },
                    "depends_on": ["step_2_calc"],
                },
            ],
        }

        req = JessycaRequest(request_id=f"req_obj2_{data_format}", user_input="procesa finanzas y haz ppt")
        resp = rt.execute_llm_plan(agent, req, plan_proposal, sandbox_dir=p_tmp)

        assert resp.success is True, f"Error: {resp.error}"
        assert resp.status == AgentExecutionState.COMPLETED
        assert len(resp.tools_executed) == 3

        # 1. Comprobar lectura
        read_out = resp.output_data["step_1_read"]
        assert read_out["rows_count"] == 4
        assert read_out["columns_count"] == 4

        # 2. Comprobar cálculos exactos
        calc_out = resp.output_data["step_2_calc"]
        stats = calc_out["summary_stats"]
        assert "Presupuesto_Total" in stats
        assert "Costo_Licencias" in stats
        # Total suma de Presupuesto_Total: 15000 + 32000 + 18500 + 14200 = 79700.0
        assert stats["Presupuesto_Total"]["sum"] == 79700.0
        # Promedio: 79700 / 4 = 19925.0
        assert stats["Presupuesto_Total"]["mean"] == 19925.0
        # Mínimo: 14200.0, Máximo: 32000.0
        assert stats["Presupuesto_Total"]["min"] == 14200.0
        assert stats["Presupuesto_Total"]["max"] == 32000.0

        # 3. Comprobar físicamente PowerPoint en disco
        assert ppt_path.exists(), f"El archivo PowerPoint {ppt_path} debe existir."
        assert ppt_path.stat().st_size > 0

        prs = pptx.Presentation(str(ppt_path))
        # Debe contener al menos 4 diapositivas: Portada, Métricas Clave, Desglose Estadístico, Conclusiones
        assert len(prs.slides) >= 4


# ── TEST 4: OBJETIVO 3 — LOCALIZAR ARCHIVOS, PROPUESTA DE ORGANIZACIÓN Y EJECUCIÓN AUTORIZADA ──

def test_objective_3_scan_preview_and_execute_authorized_organization():
    """Objetivo 3:
    1. Localizar archivos en una carpeta de origen con extensiones diversas.
    2. Proponer una organización con vista previa explicada sin modificar el disco.
    3. Ejecutar los cambios autorizados en disco creando las carpetas destino.
    4. Verificar físicamente la integridad y existencia de todos los archivos organizados."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        p_tmp = Path(tmp_dir)
        src_dir = p_tmp / "carpeta_desordenada"
        dst_dir = p_tmp / "archivos_clasificados"
        src_dir.mkdir(parents=True)

        # Crear archivos con nombres y extensiones variadas
        files_to_create = {
            "balance_2026.csv": "datos,balance,100",
            "bitacora_sistema.log": "[INFO] Sistema operativo iniciado",
            "documento_contrato.txt": "Términos del contrato de servicio",
            "esquema_datos.json": '{"modelo": "v1"}',
        }
        for name, text in files_to_create.items():
            (src_dir / name).write_text(text, encoding="utf-8")

        rt = ARERuntime(mode=ARERuntimeMode.CONTROLLED)
        agent = MockAgentForPhase13()

        plan_proposal = {
            "goal": "Localizar archivos, generar vista previa de organización y ejecutar la copia autorizada",
            "steps": [
                {
                    "step_id": "step_1_search",
                    "title": "Localizar archivos en carpeta desordenada",
                    "action_type": "files.search",
                    "parameters": {"directory": str(src_dir)},
                    "depends_on": [],
                },
                {
                    "step_id": "step_2_preview",
                    "title": "Proponer organización estructurada por extensión",
                    "action_type": "files.preview_organize",
                    "parameters": {
                        "files_list": "$step_1_search.output.found_files",
                        "target_base": str(dst_dir),
                        "rule": "by_extension",
                    },
                    "depends_on": ["step_1_search"],
                },
                {
                    "step_id": "step_3_apply",
                    "title": "Ejecutar la organización física de archivos autorizados",
                    "action_type": "files.apply_organization",
                    "parameters": {
                        "preview_mapping": "$step_2_preview.output.preview_mapping",
                        "operation": "copy",
                    },
                    "depends_on": ["step_2_preview"],
                },
                {
                    "step_id": "step_4_verify",
                    "title": "Verificar físicamente la existencia e integridad en destino",
                    "action_type": "files.verify_copies",
                    "parameters": {
                        "copies_mapping": "$step_2_preview.output.preview_mapping",
                    },
                    "depends_on": ["step_2_preview", "step_3_apply"],
                },
            ],
        }

        req = JessycaRequest(request_id="req_obj3", user_input="organiza archivos desordenados")
        resp = rt.execute_llm_plan(agent, req, plan_proposal, sandbox_dir=p_tmp)

        assert resp.success is True, f"Fallo: {resp.error}"
        assert resp.status == AgentExecutionState.COMPLETED
        assert len(resp.tools_executed) == 4

        # Comprobar el preview sin alteración prematura
        preview_out = resp.output_data["step_2_preview"]
        assert preview_out["files_count"] == 4
        assert len(preview_out["rules_explanation"]) == 4

        # Comprobar la ejecución física
        apply_out = resp.output_data["step_3_apply"]
        assert apply_out["transferred_count"] == 4
        assert apply_out["verified"] is True

        # Comprobar la verificación física
        verify_out = resp.output_data["step_4_verify"]
        assert verify_out["all_verified"] is True
        assert len(verify_out["verified_files"]) == 4
        assert len(verify_out["missing_files"]) == 0

        # Verificación directa en el sistema de archivos
        assert (dst_dir / "csv" / "balance_2026.csv").exists()
        assert (dst_dir / "log" / "bitacora_sistema.log").exists()
        assert (dst_dir / "txt" / "documento_contrato.txt").exists()
        assert (dst_dir / "json" / "esquema_datos.json").exists()


# ── TEST 5: DETENCIÓN DE DEPENDIENTES ANTE SALIDA AUSENTE O REQUISITO FALLIDO ──

def test_stop_dependents_when_output_missing_or_dependency_fails():
    """Valida que ante una salida requerida ausente o requisito no satisfecho:
    1. Se detenga inmediatamente la ejecución de dependientes.
    2. Los pasos dependientes queden registrados en el estado como BLOCKED.
    3. El estado de ejecución explique qué se completó, qué falló y qué quedó bloqueado."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        p_tmp = Path(tmp_dir)
        rt = ARERuntime(mode=ARERuntimeMode.CONTROLLED)
        agent = MockAgentForPhase13()

        # Plan donde step_2 requiere una propiedad inexistente de step_1
        plan_broken_dataflow = {
            "goal": "Probar detención de dependientes ante propiedad ausente",
            "steps": [
                {
                    "step_id": "step_1_search",
                    "title": "Búsqueda inicial",
                    "action_type": "browser.search",
                    "parameters": {"query": "consulta simple"},
                    "depends_on": [],
                },
                {
                    "step_id": "step_2_dependent",
                    "title": "Paso que requiere propiedad inexistente en la salida previa",
                    "action_type": "documents.create_word_report",
                    "parameters": {
                        "file_path": str(p_tmp / "doc_fallido.docx"),
                        "title": "Documento",
                        # 'propiedad_inexistente_xyz' no existe en la salida de browser.search
                        "sources": "$step_1_search.output.propiedad_inexistente_xyz",
                    },
                    "depends_on": ["step_1_search"],
                },
                {
                    "step_id": "step_3_downstream",
                    "title": "Paso subsiguiente que depende del paso 2",
                    "action_type": "files.read",
                    "parameters": {"file_path": str(p_tmp / "doc_fallido.docx")},
                    "depends_on": ["step_2_dependent"],
                },
            ],
        }

        req = JessycaRequest(request_id="req_broken_df", user_input="prueba dependencia rota")
        resp = rt.execute_llm_plan(agent, req, plan_broken_dataflow, sandbox_dir=p_tmp)

        assert resp.success is False
        assert resp.status == AgentExecutionState.FAILED
        assert "Fallo de dependencia de datos" in resp.response_text

        # Verificar el estado persistente y explicativo
        exec_state = resp.output_data.get("execution_state")
        assert exec_state is not None
        explanation = resp.output_data.get("explanation")
        assert explanation is not None

        # El paso 1 se completó
        assert "step_1_search" in [c["step_id"] for c in explanation["completed_steps"]]
        # El paso 2 falló por DataFlow
        assert explanation["failed_step"]["step_id"] == "step_2_dependent"
        # El paso 3 fue detenido y marcado como BLOCKED
        blocked_ids = [b["step_id"] for b in explanation["blocked_steps"]]
        assert "step_3_downstream" in blocked_ids

        # Verificar explicación en lenguaje natural
        text_explanation = explanation["summary_explanation"]
        assert "Pasos completados (1): step_1_search" in text_explanation
        assert "Paso fallido: step_2_dependent" in text_explanation
        assert "Pasos bloqueados por dependencias" in text_explanation
        assert "step_3_downstream" in text_explanation


# ── TEST 6: NO REPETIR ESCRITURAS YA REALIZADAS AL REINTENTAR (IDEMPOTENCIA DE REANUDACIÓN) ──

def test_no_repeated_writes_on_resume_and_preserves_state():
    """Valida que:
    1. Un plan que completa una escritura física (ej. Word) y falla en un paso posterior preserve el estado.
    2. Al reanudar/reintentar corrigiendo el fallo, NO vuelva a sobreescribir el archivo ya generado.
    3. Reutilice la salida previa e identifique el paso como SKIPPED/reutilizado."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        p_tmp = Path(tmp_dir)
        word_doc = p_tmp / "informe_idempotente.docx"

        rt = ARERuntime(mode=ARERuntimeMode.CONTROLLED)
        agent = MockAgentForPhase13()

        # Plan inicial donde el paso 2 falla a propósito (ej: archivo que no existe en files.read)
        initial_plan = {
            "goal": "Crear informe y luego leer archivo inexistente",
            "steps": [
                {
                    "step_id": "step_1_create_word",
                    "title": "Crear informe Word inicial",
                    "action_type": "documents.create_word_report",
                    "parameters": {
                        "file_path": str(word_doc),
                        "title": "Informe Oficial de Auditoría",
                        "sections": [{"heading": "Resumen", "paragraphs": ["Contenido inicial"]}],
                    },
                    "depends_on": [],
                },
                {
                    "step_id": "step_2_failing_read",
                    "title": "Intentar leer archivo inexistente",
                    "action_type": "files.read",
                    "parameters": {"file_path": str(p_tmp / "no_existe_archivo.txt")},
                    "depends_on": ["step_1_create_word"],
                },
            ],
        }

        req1 = JessycaRequest(request_id="req_init_fail", user_input="crea y lee")
        resp1 = rt.execute_llm_plan(agent, req1, initial_plan, sandbox_dir=p_tmp)

        assert resp1.success is False
        assert resp1.status in (AgentExecutionState.FAILED, AgentExecutionState.AWAITING_CONFIRMATION)
        assert word_doc.exists(), "El informe Word debió ser creado físicamente en el paso 1."
        initial_mtime = word_doc.stat().st_mtime_ns
        saved_state = rt.last_execution_state
        assert saved_state is not None
        assert saved_state.step_records["step_1_create_word"].status == "SUCCEEDED"
        assert saved_state.step_records["step_2_failing_read"].status == "FAILED"

        # Crear ahora el archivo faltante para corregir la causa raíz
        (p_tmp / "no_existe_archivo.txt").write_text("Archivo ahora disponible", encoding="utf-8")

        # Reanudar con resume_llm_plan pasando el previous_state
        req2 = JessycaRequest(request_id="req_resumed", user_input="reanudar plan corregido")
        resp2 = rt.resume_llm_plan(
            agent=agent,
            req=req2,
            plan_input=initial_plan,
            previous_state=saved_state,
            sandbox_dir=p_tmp,
        )

        assert resp2.success is True
        assert resp2.status == AgentExecutionState.COMPLETED

        # VERIFICACIÓN CLAVE: El archivo Word NO debió ser reescrito (mtime idéntico)
        after_mtime = word_doc.stat().st_mtime_ns
        assert after_mtime == initial_mtime, "La escritura previa no debe repetirse al reintentar."

        # El paso 1 fue reutilizado por idempotencia física
        resumed_explanation = resp2.output_data["explanation"]
        skipped_ids = [s["step_id"] for s in resumed_explanation["skipped_steps"]]
        assert "step_1_create_word" in skipped_ids


# ── TEST 7: ACLARACIONES ANTE DESTINOS ESENCIALES FALTANTES ──

def test_clarification_solicited_for_essential_missing_paths():
    """Valida que si falta una ruta de destino esencial en Word, PPTX o traslados,
    PlanValidator solicite aclaración explícita (AWAITING_CLARIFICATION)."""
    rt = ARERuntime(mode=ARERuntimeMode.CONTROLLED)
    agent = MockAgentForPhase13()

    # Plan que intenta crear un informe Word sin especificar 'file_path'
    ambiguous_word_plan = {
        "goal": "crear informe sin destino",
        "steps": [
            {
                "step_id": "step_w",
                "title": "Crear Word sin ruta",
                "action_type": "documents.create_word_report",
                "parameters": {"title": "Informe Flotante", "file_path": ""},
                "depends_on": [],
            }
        ],
    }

    req = JessycaRequest(request_id="req_clarif_word", user_input="crea un informe")
    resp = rt.execute_llm_plan(agent, req, ambiguous_word_plan)

    assert resp.status == AgentExecutionState.AWAITING_CLARIFICATION
    assert resp.requires_clarification is True
    assert "Aclaración requerida" in resp.response_text
    assert "Word" in resp.response_text


# ── TEST 8: CANCELACIÓN CONTROLADA MEDIANTE CANCELLATIONTOKEN ──

def test_cancellation_token_stops_plan_safely():
    """Valida que si se activa CancellationToken antes o durante la ejecución,
    el plan se detenga inmediatamente sin efectos residuales."""
    rt = ARERuntime(mode=ARERuntimeMode.CONTROLLED)
    agent = MockAgentForPhase13()

    token = CancellationToken()
    token.cancel()  # Cancelar antes de iniciar

    plan = {
        "goal": "plan cancelado",
        "steps": [
            {
                "step_id": "step_1",
                "title": "Búsqueda inicial",
                "action_type": "browser.search",
                "parameters": {"query": "algo"},
                "depends_on": [],
            }
        ],
    }

    req = JessycaRequest(request_id="req_cancel", user_input="cancela")
    resp = rt.execute_llm_plan(agent, req, plan, cancellation_token=token)

    assert resp.status == AgentExecutionState.INTERRUPTED
    assert "interrumpida" in resp.response_text.lower() or "cancelada" in resp.response_text.lower()
