"""Pruebas rigurosas de rutas documentales, contenido específico y garantías de seguridad.

Cubre:
1. WORD: Creación con título y 3 datos específicos suministrados; verificación de persistencia al reabrir.
2. EXCEL: Creación con cantidades, precios, fórmulas y total esperado; distinción de escritura vs recálculo.
3. POWERPOINT: Creación con múltiples diapositivas y contenido específico; verificación al reabrir.
4. Nombres y rutas con espacios, tildes y caracteres especiales.
5. Garantías de seguridad: Destino permitido (bloqueo rutas críticas), bloqueo extensiones prohibidas,
   anti-sobrescritura silenciosa y verificación física del resultado.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

import pytest
import openpyxl
import docx
import pptx

from core.local_agent.local_agent import JessycaLocalAgent
from core.local_agent.local_agent_models import InputModality, JessycaRequest
from skills.documents_skills import DocumentsCreateSkill, DocumentsReadSkill


@pytest.fixture
def agent() -> JessycaLocalAgent:
    return JessycaLocalAgent.get_instance()


@pytest.fixture
def create_skill() -> DocumentsCreateSkill:
    return DocumentsCreateSkill()


@pytest.fixture
def read_skill() -> DocumentsReadSkill:
    return DocumentsReadSkill()


# ── 1. WORD: CONTENIDO ESPECÍFICO, TILDES Y ESPACIOS ──

def test_word_creation_with_specific_content_and_accents(agent: JessycaLocalAgent, read_skill: DocumentsReadSkill):
    """Verifica que un archivo Word con título y 3 datos específicos contiene exactamente la información suministrada."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        doc_name = "Campaña de Verano y Producción 2026.docx"
        doc_path = os.path.join(tmp_dir, doc_name)

        title_esperado = "Plan Estratégico de Operaciones"
        dato_1 = "Presupuesto asignado: $45,000 USD para infraestructura"
        dato_2 = "Personal requerido: 12 ingenieros senior en telecomunicaciones"
        dato_3 = "Fecha de entrega definitiva: 15 de diciembre de 2026"

        req = JessycaRequest(
            request_id="test-word-content-01",
            user_input=f'Crea un documento word en "{doc_path}" con título "{title_esperado}" y contenido "{dato_1}"',
            modality=InputModality.TEXT,
            metadata={
                "titulo": title_esperado,
                "datos": [dato_1, dato_2, dato_3],
            },
        )
        resp = agent.interact(req)

        assert resp.success is True
        assert os.path.exists(doc_path)
        assert os.path.getsize(doc_path) > 0

        # Reabrir e inspeccionar físicamente el documento Word
        res_read = read_skill.ejecutar({"ruta": doc_path})
        assert res_read["exito"] is True
        contenido_leido = res_read["contenido"]

        assert title_esperado in contenido_leido
        assert dato_1 in contenido_leido
        assert dato_2 in contenido_leido
        assert dato_3 in contenido_leido


# ── 2. EXCEL: CANTIDADES, PRECIOS, FÓRMULAS Y DISTINCIÓN DE RECÁLCULO ──

def test_excel_creation_formulas_and_recalculation_distinction(create_skill: DocumentsCreateSkill):
    """Verifica que Excel escribe fórmulas nativas y distingue escritura de recálculo dinámico."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        xl_name = "Presupuesto de Adquisición y Logística 2026.xlsx"
        xl_path = os.path.join(tmp_dir, xl_name)

        items = [
            {"concepto": "Servidores GPU RTX 3060", "categoria": "Hardware", "cantidad": 4, "costo_unitario": 1200.0},
            {"concepto": "Licencias de Software Enterprise", "categoria": "Software", "cantidad": 10, "costo_unitario": 350.0},
            {"concepto": "Soporte Técnico Especializado", "categoria": "Servicios", "cantidad": 2, "costo_unitario": 800.0},
        ]
        # Cálculos aritméticos esperados
        subtotal_esperado = (4 * 1200.0) + (10 * 350.0) + (2 * 800.0)  # 4800 + 3500 + 1600 = 9900.0
        impuesto_esperado = subtotal_esperado * 0.18                    # 1782.0
        total_esperado = subtotal_esperado + impuesto_esperado          # 11682.0

        res_create = create_skill.ejecutar({
            "ruta": xl_path,
            "titulo": "Presupuesto de Operaciones",
            "items": items,
        })

        assert res_create["exito"] is True
        assert res_create["verified"] is True
        assert res_create["formula_written"] is True
        assert res_create["formula_escrita"] is True
        assert res_create["total_calculado_esperado"] == total_esperado
        assert res_create["resultado_esperado_calculado_independiente"] == total_esperado
        assert res_create["resultado_recalculado_motor_real"] is None
        assert res_create["recalculo_motor_real_ejecutado"] is False
        assert res_create["recalculo_motor_real_verificado"] is False
        assert res_create["calculo_verificado"] is False
        assert os.path.exists(xl_path)

        # Inspección profunda del libro con openpyxl (distinguiendo fórmula de valor evaluado)
        wb_formulas = openpyxl.load_workbook(xl_path, data_only=False)
        ws = wb_formulas.active

        # Fila 5: Servidores GPU -> E5 debe contener la fórmula =C5*D5
        assert str(ws["E5"].value).upper() == "=C5*D5"
        # Fila 6: Licencias -> E6 debe contener =C6*D6
        assert str(ws["E6"].value).upper() == "=C6*D6"
        # Fila 7: Soporte -> E7 debe contener =C7*D7
        assert str(ws["E7"].value).upper() == "=C7*D7"
        # Total Directo en E8: =SUM(E5:E7)
        assert "=SUM" in str(ws["E8"].value).upper()

        wb_formulas.close()


# ── 3. POWERPOINT: MÚLTIPLES DIAPOSITIVAS Y CONTENIDO COMPROBADO ──

def test_powerpoint_multiple_slides_with_accents_and_spaces(create_skill: DocumentsCreateSkill, read_skill: DocumentsReadSkill):
    """Verifica que PowerPoint genera múltiples diapositivas con contenido específico y las reabre sin fallos."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        ppt_name = "Presentación Estratégica de Innovación Tecnológica.pptx"
        ppt_path = os.path.join(tmp_dir, ppt_name)

        slides = [
            {
                "title": "Visión General del Proyecto",
                "subtitle": "Estrategia 2026-2027",
                "category_label": "ESTRATEGIA",
            },
            {
                "title": "Hitos de Producción y Entrega",
                "bullets": ["Fase 1 completada al 100%", "Fase 2 en despliegue controlado", "Certificación de calidad ISO"],
                "category_label": "OPERACIONES",
            },
            {
                "title": "Métricas de Calidad y Rendimiento",
                "bullets": ["Disponibilidad de servicio 99.9%", "Latencia promedio menor a 250ms", "Cero vulnerabilidades críticas"],
                "category_label": "MÉTRICAS",
            },
        ]

        res_create = create_skill.ejecutar({
            "ruta": ppt_path,
            "titulo": "Innovación Tecnológica",
            "slides": slides,
            "presenter": "Dirección de Ingeniería",
        })

        assert res_create["exito"] is True
        assert res_create["verified"] is True
        assert res_create["slides_count"] == 3
        assert os.path.exists(ppt_path)
        assert os.path.getsize(ppt_path) > 0

        # Reabrir e inspeccionar contenido estructurado
        res_read = read_skill.ejecutar({"ruta": ppt_path})
        assert res_read["exito"] is True
        texto = res_read["contenido"]

        assert "Visión General del Proyecto" in texto
        assert "Hitos de Producción y Entrega" in texto
        assert "Fase 1 completada al 100%" in texto
        assert "Métricas de Calidad y Rendimiento" in texto
        assert "Disponibilidad de servicio 99.9%" in texto


# ── 4. GARANTÍAS DE SEGURIDAD Y PROTECCIÓN ──

def test_security_guarantee_critical_path_blocked(create_skill: DocumentsCreateSkill):
    """Garantía de Destino Permitido: Prohíbe escribir en rutas sensibles del sistema operativo."""
    ruta_critica = r"C:\Windows\System32\documento_no_permitido.docx"
    res = create_skill.ejecutar({
        "ruta": ruta_critica,
        "titulo": "Intento Crítico",
    })

    assert res["exito"] is False
    assert res["verified"] is False
    assert "Acceso denegado" in res["mensaje"] or "crítica" in res["mensaje"]


def test_security_guarantee_forbidden_extension_blocked(create_skill: DocumentsCreateSkill):
    """Garantía de Extensión Segura: Prohíbe generar binarios o scripts ejecutables mediante la skill de documentos."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        ruta_ejecutable = os.path.join(tmp_dir, "script_malicioso.bat")
        res = create_skill.ejecutar({
            "ruta": ruta_ejecutable,
            "titulo": "Intento de Script",
            "contenido": "echo hack",
        })

        assert res["exito"] is False
        assert res["verified"] is False
        assert "peligrosa bloqueada" in res["mensaje"] or "Seguridad" in res["mensaje"]


def test_security_guarantee_anti_silent_overwrite(create_skill: DocumentsCreateSkill):
    """Garantía Anti-sobrescritura Silenciosa: No sobrescribe archivos existentes a menos que se confirme expresamente."""
    from core.jde.confirmation import get_confirmation_manager
    conf_mgr = get_confirmation_manager()
    conf_mgr.reset()

    session_id = "test-session-doc-001"

    with tempfile.TemporaryDirectory() as tmp_dir:
        archivo = os.path.join(tmp_dir, "reporte_existente.docx")
        destino_norm = os.path.abspath(archivo)

        # 1. Crear por primera vez
        res1 = create_skill.ejecutar({
            "ruta": archivo,
            "titulo": "Versión 1",
            "contenido": "Contenido original",
            "session_id": session_id,
        })
        assert res1["exito"] is True
        assert res1["archivo_creado"] is True
        assert res1["contenido_verificado"] is True

        # 2. Intentar crear nuevamente sin confirmación (debe fallar y registrar confirmación pendiente)
        res2 = create_skill.ejecutar({
            "ruta": archivo,
            "titulo": "Versión 2",
            "contenido": "Sobrescritura no autorizada",
            "sobrescribir": False,
            "session_id": session_id,
        })
        assert res2["exito"] is False
        assert res2["verified"] is False
        assert res2.get("conflicto_destino") is True
        assert res2.get("requiere_confirmacion") is True
        assert "ya existe" in res2["mensaje"]
        assert "anti-sobrescritura" in res2["mensaje"]

        # 3. Intentar crear con sobrescribir=True pero SIN aprobación humana en JDE (debe ser RECHAZADO)
        res3 = create_skill.ejecutar({
            "ruta": archivo,
            "titulo": "Versión Inyectada",
            "contenido": "Intento de bypass de LLM o cliente",
            "sobrescribir": True,
            "session_id": session_id,
        })
        assert res3["exito"] is False
        assert res3["verified"] is False
        assert res3.get("sobrescritura_no_autorizada") is True
        assert "no constituye autorización" in res3["mensaje"]

        # 4. Autorizar legítimamente en JDE para esta sesión, operación y destino exacto
        aprobado = conf_mgr.confirm(
            session_id=session_id,
            action="documents.create",
            parameters={"ruta": destino_norm},
        )
        assert aprobado is True

        # 5. Intentar crear con sobrescribir=True tras confirmación vinculada (debe autorizarse)
        res4 = create_skill.ejecutar({
            "ruta": archivo,
            "titulo": "Versión 2 Autorizada",
            "contenido": "Contenido reemplazado legítimamente",
            "sobrescribir": True,
            "session_id": session_id,
        })
        assert res4["exito"] is True
        assert res4["verified"] is True
        assert res4["archivo_creado"] is True
        assert res4["contenido_verificado"] is True


def test_rejection_of_unauthorized_overwrite_injection_and_session_binding(create_skill: DocumentsCreateSkill):
    """Regresión: Demuestra que un LLM o cliente NO puede convertir sobrescribir=True en autorización.

    Verifica la vinculación estricta a:
    - Sesión (sesión B no puede consumir autorización de sesión A)
    - Destino (autorización para archivo X no autoriza archivo Y)
    - Monouso (una confirmación gastada no puede reutilizarse)
    """
    from core.jde.confirmation import get_confirmation_manager
    conf_mgr = get_confirmation_manager()
    conf_mgr.reset()

    session_a = "session-human-owner-A"
    session_b = "session-adversary-or-llm-B"

    with tempfile.TemporaryDirectory() as tmp_dir:
        doc_x = os.path.join(tmp_dir, "archivo_x.docx")
        doc_y = os.path.join(tmp_dir, "archivo_y.docx")
        doc_x_norm = os.path.abspath(doc_x)
        doc_y_norm = os.path.abspath(doc_y)

        # Crear archivos iniciales
        assert create_skill.ejecutar({"ruta": doc_x, "titulo": "X1", "contenido": "Data X"})["exito"] is True
        assert create_skill.ejecutar({"ruta": doc_y, "titulo": "Y1", "contenido": "Data Y"})["exito"] is True

        # Paso 1: Sesión A registra pendiente para doc_x
        res_pend = create_skill.ejecutar({
            "ruta": doc_x,
            "titulo": "X2",
            "contenido": "Nuevo X",
            "sobrescribir": False,
            "session_id": session_a,
        })
        assert res_pend.get("requiere_confirmacion") is True

        # Paso 2: Sesión A aprueba confirmación para doc_x
        assert conf_mgr.confirm(session_id=session_a, action="documents.create", parameters={"ruta": doc_x_norm}) is True

        # Paso 3: Sesión B (LLM o atacante) intenta sobrescribir doc_x con sobrescribir=True -> RECHAZADO
        res_b = create_skill.ejecutar({
            "ruta": doc_x,
            "titulo": "Hack X",
            "contenido": "Inyección",
            "sobrescribir": True,
            "session_id": session_b,
        })
        assert res_b["exito"] is False
        assert res_b.get("sobrescritura_no_autorizada") is True

        # Paso 4: Sesión A intenta usar esa confirmación para sobrescribir doc_y (destino distinto) -> RECHAZADO
        res_wrong_dest = create_skill.ejecutar({
            "ruta": doc_y,
            "titulo": "Hack Y",
            "contenido": "Desvío de destino",
            "sobrescribir": True,
            "session_id": session_a,
        })
        assert res_wrong_dest["exito"] is False
        assert res_wrong_dest.get("sobrescritura_no_autorizada") is True

        # Paso 5: Sesión A sobrescribe doc_x legítimamente -> ACEPTADO y CONSUMIDO
        res_ok = create_skill.ejecutar({
            "ruta": doc_x,
            "titulo": "X2 Autorizado",
            "contenido": "Actualización legítima",
            "sobrescribir": True,
            "session_id": session_a,
        })
        assert res_ok["exito"] is True
        assert res_ok["verified"] is True

        # Paso 6: Reintento posterior en sesión A (monouso) -> RECHAZADO porque ya se consumió
        res_reused = create_skill.ejecutar({
            "ruta": doc_x,
            "titulo": "X3 Sin Reconfirmar",
            "contenido": "Reintento de reuso",
            "sobrescribir": True,
            "session_id": session_a,
        })
        assert res_reused["exito"] is False
        assert res_reused.get("sobrescritura_no_autorizada") is True
