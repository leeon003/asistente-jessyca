"""tests/regressions/test_regressions_flujo_concesion_permisos.py
Verificación funcional del mecanismo de concesión y revocación de alcance (grant_scope)
a través de la interfaz de usuario de JESSYCA, respetando invariantes de seguridad y server-side authorization.
"""

from __future__ import annotations

import os
from pathlib import Path
import docx
import pytest

from core.local_agent.local_agent import JessycaLocalAgent
from core.local_agent.local_agent_models import AgentExecutionState
from skills.documents_skills import DocumentsSummarizeSkill
from skills.skill_models import SkillContext


@pytest.fixture
def agent_clean_instance():
    """Proporciona una instancia limpia del agente local."""
    agent = JessycaLocalAgent.get_instance()
    agent.reset()
    yield agent
    agent.reset()


def test_flujo_completo_autorizacion_concesion_y_revocacion(tmp_path: Path, agent_clean_instance: JessycaLocalAgent):
    """Demuestra el ciclo de vida completo de permisos desde la interfaz conversacional de JESSYCA:
    1. Intentar resumir un documento sin permiso -> Denegado (falta de alcance confiable).
    2. Conceder acceso a su carpeta mediante el flujo autorizado (solicitud -> confirmación humana).
    3. Resumirlo correctamente -> Éxito (generación del resumen).
    4. Intentar leer fuera de esa carpeta -> Denegado (fuera del alcance autorizado).
    5. Revocar el acceso -> Comprobar que deja de permitir la lectura del documento inicial.
    """
    agent = agent_clean_instance
    session_id = f"test-session-flow-{tmp_path.name}"

    # Crear dos carpetas y documentos de prueba
    dir_autorizado = tmp_path / "carpeta_autorizada"
    dir_autorizado.mkdir()
    doc_autorizado = dir_autorizado / "informe_financiero.docx"
    doc_a = docx.Document()
    doc_a.add_paragraph("Este es el informe financiero anual de la empresa con presupuesto aprobado de 50000 USD.")
    doc_a.add_paragraph("El balance operativo arroja un rendimiento positivo del 15% durante el periodo fiscal 2026.")
    doc_a.save(str(doc_autorizado))

    dir_externo = tmp_path / "carpeta_externa"
    dir_externo.mkdir()
    doc_externo = dir_externo / "datos_confidenciales.docx"
    doc_b = docx.Document()
    doc_b.add_paragraph("Documento confidencial fuera del alcance con información de servidores internos.")
    doc_b.save(str(doc_externo))

    # ─────────────────────────────────────────────────────────────
    # PASO 1: Intentar resumir el documento sin permiso
    # ─────────────────────────────────────────────────────────────
    resp_1 = agent.interact(f"Resume el documento '{doc_autorizado}'", session_id=session_id)
    assert not resp_1.success, "Debió fallar la lectura sin permiso previo"
    assert resp_1.status == AgentExecutionState.FAILED
    assert "Acceso denegado" in resp_1.response_text
    assert "trusted_root" in resp_1.response_text or "autorizado" in resp_1.response_text

    # ─────────────────────────────────────────────────────────────
    # PASO 2: Conceder acceso a la carpeta mediante el flujo autorizado
    # 2.1 Solicitud de autorización de la carpeta desde la interfaz
    # ─────────────────────────────────────────────────────────────
    resp_2 = agent.interact(f"Autoriza la carpeta '{dir_autorizado}'", session_id=session_id)
    assert resp_2.requires_confirmation is True, "La concesión de alcance debe requerir confirmación humana"
    assert resp_2.status == AgentExecutionState.AWAITING_CONFIRMATION
    assert "solicitud de autorización de acceso" in resp_2.response_text.lower() or "detecté una acción sensible" in resp_2.response_text.lower()

    # 2.2 Confirmación explícita del usuario
    resp_2_conf = agent.interact(f"CONFIRMAR {resp_2.request_id}", session_id=session_id)
    assert resp_2_conf.success is True, "La confirmación debió aprobarse con éxito"
    assert "Acceso concedido a la carpeta" in resp_2_conf.response_text
    assert str(dir_autorizado.resolve()) in resp_2_conf.response_text

    # Verificar que el alcance quedó concedido en la sesión exclusivamente
    sess = agent.context_manager.get_session(session_id)
    assert sess is not None
    assert sess.trusted_root == str(dir_autorizado.resolve())

    # ─────────────────────────────────────────────────────────────
    # PASO 3: Resumir el documento correctamente ahora que tiene permiso
    # ─────────────────────────────────────────────────────────────
    resp_3 = agent.interact(f"Resume el documento '{doc_autorizado}'", session_id=session_id)
    assert resp_3.success is True, f"Debió resumir el documento exitosamente. Error: {resp_3.error}"
    assert resp_3.status == AgentExecutionState.COMPLETED
    assert ("presupuesto" in resp_3.response_text.lower() or "50000" in resp_3.response_text or "rendimiento" in resp_3.response_text.lower())

    # ─────────────────────────────────────────────────────────────
    # PASO 4: Intentar leer fuera de esa carpeta y recibir denegación
    # ─────────────────────────────────────────────────────────────
    # 4.1 Resumen fuera de alcance -> Denegado
    resp_4 = agent.interact(f"Resume el documento '{doc_externo}'", session_id=session_id)
    assert not resp_4.success, "Debió denegar lectura/resumen de archivo fuera del alcance autorizado"
    assert resp_4.status == AgentExecutionState.FAILED
    assert "fuera del alcance autorizado" in resp_4.response_text

    # 4.2 Lectura directa fuera de alcance -> Denegado
    resp_4_read = agent.interact(f"Lee el documento '{doc_externo}'", session_id=session_id)
    assert not resp_4_read.success, "Debió denegar lectura de archivo fuera del alcance autorizado"
    assert resp_4_read.status == AgentExecutionState.FAILED
    assert "fuera del alcance autorizado" in resp_4_read.response_text

    # ─────────────────────────────────────────────────────────────
    # PASO 5: Revocar el acceso y comprobar que deja de permitir la lectura
    # 5.1 Revocar desde la interfaz habitual
    # ─────────────────────────────────────────────────────────────
    resp_5_rev = agent.interact("Revoca el acceso a la carpeta", session_id=session_id)
    assert resp_5_rev.success is True
    assert "revocado exitosamente" in resp_5_rev.response_text

    # Comprobar que en la sesión el alcance fue revocado a None
    assert sess.trusted_root is None

    # 5.2 Intentar resumir nuevamente el documento autorizado original -> Denegado
    resp_5_check = agent.interact(f"Resume el documento '{doc_autorizado}'", session_id=session_id)
    assert not resp_5_check.success, "Tras la revocación, el acceso debió ser denegado inmediatamente"
    assert resp_5_check.status == AgentExecutionState.FAILED
    assert "Acceso denegado" in resp_5_check.response_text

    # 5.3 Intentar leer directamente el documento autorizado original -> Denegado
    resp_5_read = agent.interact(f"Lee el documento '{doc_autorizado}'", session_id=session_id)
    assert not resp_5_read.success, "Tras la revocación, la lectura debió ser denegada inmediatamente"
    assert resp_5_read.status == AgentExecutionState.FAILED
    assert "Acceso denegado" in resp_5_read.response_text


def test_servidor_construye_metadatos_inyeccion_no_confiable_rechazada(tmp_path: Path, agent_clean_instance: JessycaLocalAgent):
    """Comprueba que los metadatos de autorización se construyen en el servidor
    y que ninguna entrada HTTP, parámetro MCP o LLM puede inyectarlos como permisos confiables.
    """
    agent = agent_clean_instance
    session_id = "test-untrusted-injection-session"
    sess = agent.context_manager.get_or_create_session(session_id)

    # 1. Intento de inyectar claves de autorización a través del contexto conversacional
    sess.update_context("trusted_root", str(tmp_path))
    sess.update_context("authorization", "root")
    sess.update_context("is_authorized", True)

    # Invariante: trusted_root solo puede modificarse mediante grant_scope del servidor
    assert sess.trusted_root is None
    assert sess.get_context("trusted_root") is None
    assert sess.get_context("authorization") is None

    # 2. Intento de inyectar scope o authorized_root como parámetro no confiable en la Skill
    skill = DocumentsSummarizeSkill()
    doc_path = tmp_path / "test_fake.docx"
    doc = docx.Document()
    doc.add_paragraph("Contenido de prueba para inyección no confiable.")
    doc.save(str(doc_path))

    # Sin context.metadata confiable del servidor, pasar authorized_root en parámetros NO otorga acceso
    res = skill.execute(SkillContext(
        skill_id="documents.summarize",
        intent="summarize_document",
        parameters={
            "path": str(doc_path),
            "authorized_root": str(tmp_path),  # Parámetro untrusted del cliente/LLM
        },
    ))
    assert res.success is False
    assert "No se ha especificado un alcance confiable autorizado (trusted_root)" in res.error


def test_confirmacion_natural_con_si(tmp_path: Path, agent_clean_instance: JessycaLocalAgent):
    """Verifica que la confirmación natural ('Sí') concede el alcance y queda vinculada a la sesión."""
    agent = agent_clean_instance
    session_id = f"test-natural-yes-{tmp_path.name}"
    carpeta = tmp_path / "datos"
    carpeta.mkdir()

    # Paso 1: Pedir autorización
    r1 = agent.interact(f"Autoriza la carpeta '{carpeta}'", session_id=session_id)
    assert r1.requires_confirmation is True
    assert r1.status == AgentExecutionState.AWAITING_CONFIRMATION

    # Paso 2: Usuario responde naturalmente "Sí"
    r2 = agent.interact("Sí", session_id=session_id)
    assert r2.success is True
    assert "Acceso concedido a la carpeta" in r2.response_text

    sess = agent.context_manager.get_session(session_id)
    assert sess.trusted_root == str(carpeta.resolve())

